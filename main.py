import streamlit as st
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go

from datetime import datetime
from zoneinfo import ZoneInfo
from io import BytesIO
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from streamlit_sortables import sort_items
from google import genai
from google.genai import types
from pydantic import BaseModel
from services.google_drive import upload_screenshot

from services.google_sheets import (
    get_google_sheet_data,
    append_pengaduan_row,
    get_pengaduan_data,
    append_pemutakhiran_row,
    get_pemutakhiran_data,
    append_faq_row,
    get_faq_data,
    get_skd_data,
    upsert_skd_row,
)

# --- CONFIG HALAMAN ---
st.set_page_config(
    page_title="JALA-SERIBU",
    page_icon="assets/logo_jala_seribu.png",
    layout="wide"
)

# --- CUSTOM CSS ---
st.markdown(
    """
    <style>
        /* Hide Streamlit Toolbar */
        [data-testid="stToolbar"] {
            display: none !important;
        }

        [data-testid="stDecoration"] {
            display: none !important;
        }

        [data-testid="stStatusWidget"] {
            display: none !important;
        }

        header[data-testid="stHeader"] {
            display: none !important;
        }

        .block-container {
            padding-top: 0.2rem !important;
            padding-bottom: 1rem !important;

            padding-left: 2rem !important;
            padding-right: 2rem !important;

            max-width: 100% !important;
        }

        /* =========================
        HEADER JALA-SERIBU
        ========================= */

        .jala-brand {
            display: flex;
            align-items: center;
            gap: 12px;

            height: auto !important;
            min-height: 0 !important;

            padding: 6px 0 4px 0;
            margin: 0;

            overflow: visible !important;
            
        }

        .jala-logo {
            width: 44px;
            height: 44px;
            object-fit: contain;
            flex-shrink: 0;
        }

        .jala-brand-text {
            display: flex;
            flex-direction: column;
            justify-content: center;

            height: auto !important;
            overflow: visible !important;
            gap: 2px;
        }

        .jala-title {
            color: #002B6A;
            font-size: 25px;
            font-weight: 800;
            letter-spacing: -0.8px;

            line-height: 0.7 !important;

            margin-top: 5px !important;
            padding: 0;

            overflow: visible !important;
        }

        .jala-subtitle {
            color: #8A8F98;
            font-size: 14px;
            line-height: 1.25;
            margin-top: 5px;
        }

        /* =========================================
        TOP NAVIGATION JALA-SERIBU
        ========================================= */

        /* Anchor posisi navbar */
        div[data-testid="stTabs"] {
            position: relative !important;
        }

        /* KAPSUL BESAR NAVBAR */
        div[data-testid="stTabs"] div[data-baseweb="tab-list"] {
            position: absolute !important;
            right: 0 !important;
            top: -48px !important;

            width: auto !important;

            display: flex !important;
            justify-content: flex-end !important;
            align-items: center !important;

            gap: 4px !important;

            padding: 3px !important;

            background: linear-gradient(
                135deg,
                #F8FBFF 0%,
                #F1F6FC 100%
            ) !important;

            border: 1px solid #DCE6F2 !important;
            border-radius: 999px !important;

            box-shadow: 0 2px 8px rgba(0, 43, 106, 0.05) !important;

            z-index: 10 !important;
        }

        /* MENU BIASA */
        div[data-testid="stTabs"] button[data-baseweb="tab"] {
            height: 32px !important;

            padding: 0 16px !important;

            border-radius: 999px !important;
            border: none !important;

            background: transparent !important;

            color: #475569 !important;
            font-size: 14px !important;
            font-weight: 500 !important;

            transition: all 0.15s ease !important;
        }

        /* HOVER */
        div[data-testid="stTabs"] button[data-baseweb="tab"]:hover {
            background: rgba(255, 255, 255, 0.65) !important;
            color: #002B6A !important;
        }

        /* MENU AKTIF */
        div[data-testid="stTabs"]
        button[data-baseweb="tab"][aria-selected="true"] {
            background: linear-gradient(
                135deg,
                #FFFFFF 0%,
                #EAF3FF 55%,
                #DCEBFF 100%
            ) !important;

            color: #002B6A !important;
            font-weight: 700 !important;

            box-shadow:
                0 2px 5px rgba(0, 43, 106, 0.10),
                inset 0 0 0 1px rgba(0, 43, 106, 0.06) !important;
        }

        /* HILANGKAN GARIS AKTIF BAWAAN */
        div[data-testid="stTabs"] div[data-baseweb="tab-highlight"],
        div[data-testid="stTabs"] div[data-baseweb="tab-border"] {
            display: none !important;
        }

        /* Menghilangkan jarak bawah pada judul utama (h1) */
        div[data-testid="stHeadingWithActionElements"] h1,
        .stMarkdown h1 {
            color: #002B6A !important;
            margin-bottom: 0px !important;
            padding-bottom: 5px !important;
        }

        /* JUDUL UTAMA SETIAP HALAMAN */
        .jala-page-title {
            color: #16477C !important;
            font-size: 30px;
            font-weight: 600;
            line-height: 1.2;
            margin: 8px 0 15px 0;
            letter-spacing: -0.3px;
        }

        /* Caption khusus baris update + periode + download */
        div[data-testid="stHorizontalBlock"]:has(
            div[data-testid="stSelectbox"]
        ) div[data-testid="stCaptionContainer"] {
            margin-top: -5px !important;
            margin-bottom: -20px !important;
            padding-top: 0px !important;
        }

        /* Mengubah wadah st.metric jadi card */
        div[data-testid="stMetric"] {
            background: linear-gradient(135deg, #F8FAFC 0%, #DBEAFE 100%);
            border: 1px solid #CBD5E1;
            border-top: 4px solid #002B6A; 
            margin-bottom: 0px !important;
            padding: 16px 20px;
            border-radius: 12px;
            box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.05);
            transition: transform 0.2s ease, box-shadow 0.2s ease;
            min-height: 130px;
        }
        
        /* Efek Hover */
        div[data-testid="stMetric"]:hover {
            transform: translateY(-3px);
            box-shadow: 0 10px 15px -3px rgba(0, 0, 0, 0.1);
        }

        /* Judul metrik */
        div[data-testid="stMetricLabel"] > div {
            font-size: 14px !important;
            color: #475569 !important;
            font-weight: 600 !important;
        }

        /* Warna angka metrik */
        div[data-testid="stMetricValue"] > div {
            font-size: 28px !important;
            font-weight: 800 !important;
            color: #0F172A !important;
        }

        /* Aturan dasar badge delta */
        div[data-testid="stMetricDelta"] {
            padding: 2px 8px !important;
            border-radius: 20px !important;
            width: fit-content !important;
            font-weight: 700 !important;
        }

        /* Badge hijau (panah naik) */
        div[data-testid="stMetricDelta"]:has([data-testid="stMetricDeltaIcon-Up"]) {
            background-color: #DCFCE7 !important;
            color: #15803D !important;
        }
        div[data-testid="stMetricDelta"] [data-testid="stMetricDeltaIcon-Up"] {
            color: #16A34A !important;
        }

        /* Badge merah (panah turun) */
        div[data-testid="stMetricDelta"]:has([data-testid="stMetricDeltaIcon-Down"]) {
            background-color: #FEE2E2 !important;
            color: #B91C1C !important;
        }
        div[data-testid="stMetricDelta"] [data-testid="stMetricDeltaIcon-Down"] {
            color: #DC2626 !important;
        }

        /* Naikin garis pembatas */
        hr {
            margin-top: 5px !important;    /* Atur jarak atas garis (makin kecil makin naik) */
            margin-bottom: 5px !important; /* Jarak dari garis ke judul chart bawah */
        }

        div[data-testid="stVerticalBlockBorderWrapper"] > div {
            background: linear-gradient(135deg, #F8FAFC 0%, #EFF6FF 100%) !important;
            border: 1px solid #CBD5E1 !important;
            border-top: 4px solid #002B6A !important;
            border-radius: 12px !important;
            padding: 16px !important;
            box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.05) !important;
        }

        div[data-testid="stDownloadButton"] button {
            width: 105px !important;
            min-width: 105px !important;
            max-width: 105px !important;

            height: 30px !important;
            min-height: 30px !important;

            padding: 3px 7px !important;
            white-space: nowrap !important;
        }

        /* Ukuran tulisan */
        div[data-testid="stDownloadButton"] button p {
            font-size: 13px !important;
        }

        div[data-testid="stDownloadButton"] button:hover {
            background: #F0F7FF !important;
            border-color: #64B5F6 !important;
            color: #0D47A1 !important;
        }

        div[data-testid="stSelectbox"] div[data-baseweb="select"] > div {
            min-height: 30px !important;
            height: 30px !important;
            font-size: 13px !important;
            padding-top: 0px !important;
            padding-bottom: 0px !important;
        }

        /* Wrapper isi selectbox: teks + ikon */
        div[data-testid="stSelectbox"]
        div[data-baseweb="select"] > div > div {
            min-height: 30px !important;
            height: 30px !important;

            display: flex !important;
            align-items: center !important;

            padding-top: 1px !important;
            padding-bottom: 0px !important;
            box-sizing: border-box !important;
        }

        /* Teks pilihan */
        div[data-testid="stSelectbox"]
        div[data-baseweb="select"] div[value] {
            line-height: 1 !important;
            padding-top: 8.5px !important;
            margin: 0 !important;
            transform: none !important;
        }

        /* =========================================
        UPDATE + PERIODE + DOWNLOAD AUTO RESPONSIVE
        ========================================= */

        div[data-testid="stHorizontalBlock"]:has(
            div[data-testid="stSelectbox"]
        ) {
            display: flex !important;
            flex-wrap: wrap !important;
            align-items: center !important;

            column-gap: 12px !important;
            row-gap: 6px !important;

            margin-top: -11px !important;
            margin-bottom: 1px !important;
        }

        /* Update fleksibel */
        div[data-testid="stHorizontalBlock"]:has(
            div[data-testid="stSelectbox"]
        ) > div:nth-child(1) {
            flex: 1 1 420px !important;
            width: auto !important;
            min-width: 280px !important;
        }

        /* Periode fleksibel */
        div[data-testid="stHorizontalBlock"]:has(
            div[data-testid="stSelectbox"]
        ) > div:nth-child(2) {
            flex: 0 1 230px !important;
            width: auto !important;
            min-width: 180px !important;
            max-width: 230px !important;
        }

        /* Download tetap */
        div[data-testid="stHorizontalBlock"]:has(
            div[data-testid="stSelectbox"]
        ) > div:nth-child(3) {
            flex: 0 0 105px !important;
            width: 105px !important;
            min-width: 105px !important;
            max-width: 105px !important;
            margin: 0 !important;
        }

        /* Selectbox ikut ruang yang tersedia */
        div[data-testid="stSelectbox"] {
            width: 100% !important;
            max-width: 230px !important;
        }

        /* Tombol download */
        div[data-testid="stDownloadButton"] {
            display: flex !important;
            justify-content: flex-start !important;
        }

        .jala-footer {
            margin-top: 50px;
            padding: 15px 0 6px 0;
            border-top: 1px solid #E7EDF5;

            text-align: center;
            color: #8A94A3;
            font-size: 12px;
            font-weight: 400;
        }

        .footer-bps {
            color: #6F91B5;
            font-weight: 500;
        }

        /* =========================================
        RESPONSIVE - MOBILE
        ========================================= */
        @media (max-width: 700px) {

            /*
            Periode + download mengikuti
            layout auto-responsive di atas.
            */

            /* Metric cards vertikal */
            div[data-testid="stHorizontalBlock"]:has(
                div[data-testid="stMetric"]
            ) {
                flex-direction: column !important;
                gap: 10px !important;
            }

            div[data-testid="stHorizontalBlock"]:has(
                div[data-testid="stMetric"]
            ) > div {
                width: 100% !important;
                min-width: 100% !important;
                flex: 0 0 auto !important;
            }

            /* Chart vertikal */
            div[data-testid="stHorizontalBlock"]:has(
                div[data-testid="stPlotlyChart"]
            ) {
                flex-direction: column !important;
                gap: 10px !important;
            }

            div[data-testid="stHorizontalBlock"]:has(
                div[data-testid="stPlotlyChart"]
            ) > div {
                width: 100% !important;
                min-width: 100% !important;
                flex: 0 0 auto !important;
            }
        }


        /* =========================================
        MOBILE SANGAT KECIL
        ========================================= */
        @media (max-width: 380px) {

            /* Kalau benar-benar sempit, baru kontrol ditumpuk */
            div[data-testid="stHorizontalBlock"]:has(
                div[data-testid="stSelectbox"]
            ) > div:nth-child(2),
            div[data-testid="stHorizontalBlock"]:has(
                div[data-testid="stSelectbox"]
            ) > div:nth-child(3) {
                flex: 0 0 100% !important;
                width: 100% !important;
                max-width: 100% !important;
                min-width: 0 !important;
            }

            div[data-testid="stSelectbox"] {
                width: 100% !important;
                max-width: 100% !important;
            }
        }    

        @media (max-width: 1050px) {

            div[data-testid="stTabs"] div[data-baseweb="tab-list"] {
                position: static !important;

                width: fit-content !important;
                max-width: 100% !important;

                justify-content: flex-start !important;

                overflow-x: auto !important;
                white-space: nowrap !important;

                margin-top: 4px !important;
                margin-bottom: 8px !important;
            }

            .jala-brand {
                margin-bottom: 4px !important;
            }
        }

    </style>
""",
    unsafe_allow_html=True,
)

# --- HEADER/TITLE ---
import base64

with open("assets/logo_jala_seribu.png", "rb") as f:
    logo_base64 = base64.b64encode(f.read()).decode()

st.markdown(
    f"""<div class="jala-brand">
<img src="data:image/png;base64,{logo_base64}" class="jala-logo">
<div class="jala-brand-text">
<div class="jala-title">JALA-SERIBU</div>
<div class="jala-subtitle">Jaringan Agregasi Layanan dan Akuntabilitas BPS Kabupaten Kepulauan Seribu</div>
</div>
</div>""",
    unsafe_allow_html=True
)

# --- GLOBAL CONSTANTA ---
kategori_jenis_tamu = [
    "Masyarakat Umum/Mahasiswa",
    "Mitra Statistik",
    "Instansi/Dinas",
    "Lainnya"
]

kategori_keperluan = [
    "Konsultasi Data",
    "Tugas Mitra",
    "Kepentingan Dinas",
    "Bertemu Orang",
    "Lainnya"
]

# --- HELPER FUNCTIONS ---
def buat_excel_overview(
    df_ringkasan,
    data_kunjungan,
    data_jenis,
    df_mentah,
    rekap_bulanan=None, 
    nama_periode=None
):
    output = BytesIO()

    if nama_periode:
        bagian_periode = nama_periode.split("_")

        if bagian_periode[0] == "Triwulan":
            nomor_triwulan = int(bagian_periode[1])
            tahun_periode = bagian_periode[2]

            judul_periode = (
                f"TRIWULAN {romawi[nomor_triwulan]} "
                f"TAHUN {tahun_periode}"
            )

        else:
            judul_periode = nama_periode.replace("_", " ").upper()

    else:
        judul_periode = ""

            # =========================
            # WARNA & STYLE DASAR
            # =========================
    biru = "1F4E78"
    putih = "FFFFFF"

    garis_tipis = Side(
        style="thin",
        color="000000"
    )

    border_tabel = Border(
        left=garis_tipis,
        right=garis_tipis,
        top=garis_tipis,
        bottom=garis_tipis
    )

    fill_header = PatternFill(
        fill_type="solid",
        fgColor=biru
    )

    font_header = Font(
        bold=True,
        color=putih
    )

    align_center = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

    with pd.ExcelWriter(output, engine="openpyxl") as writer:

        df_ringkasan.to_excel(
            writer,
            index=False,
            sheet_name="Ringkasan"
        )

        ws_ringkasan = writer.sheets["Ringkasan"]

        # Sisipkan 2 baris di atas untuk judul
        ws_ringkasan.insert_rows(1, amount=2)

        # Judul periode
        ws_ringkasan.merge_cells("A1:B1")
        ws_ringkasan["A1"] = judul_periode
        ws_ringkasan["A1"].font = Font(
            bold=True,
            size=14
        )
        ws_ringkasan["A1"].alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        # Header sekarang pindah ke baris 3
        for cell in ws_ringkasan[3]:
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = align_center
            cell.border = border_tabel

        # Isi tabel
        for row in ws_ringkasan.iter_rows(
            min_row=4,
            max_row=ws_ringkasan.max_row,
            min_col=1,
            max_col=2
        ):
            for cell in row:
                cell.border = border_tabel
                cell.alignment = Alignment(
                    vertical="center"
                )

        # Kolom Nilai dibuat center
        for row in range(4, ws_ringkasan.max_row + 1):
            ws_ringkasan.cell(
                row=row,
                column=2
            ).alignment = align_center

        # Lebar kolom
        ws_ringkasan.column_dimensions["A"].width = 32
        ws_ringkasan.column_dimensions["B"].width = 28

        data_kunjungan.to_excel(
            writer,
            index=False,
            sheet_name="Keperluan Kunjungan"
        )

        ws_keperluan = writer.sheets["Keperluan Kunjungan"]

        # Sisipkan 2 baris untuk judul
        ws_keperluan.insert_rows(1, amount=2)

        # Judul periode
        ws_keperluan.merge_cells("A1:B1")
        ws_keperluan["A1"] = judul_periode
        ws_keperluan["A1"].font = Font(
            bold=True,
            size=14
        )
        ws_keperluan["A1"].alignment = align_center

        # Header
        for cell in ws_keperluan[3]:
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = align_center
            cell.border = border_tabel

        # Isi tabel
        for row in ws_keperluan.iter_rows(
            min_row=4,
            max_row=ws_keperluan.max_row,
            min_col=1,
            max_col=2
        ):
            for cell in row:
                cell.border = border_tabel
                cell.alignment = Alignment(
                    vertical="center"
                )

        # Kolom jumlah dibuat center
        for row in range(4, ws_keperluan.max_row + 1):
            ws_keperluan.cell(
                row=row,
                column=2
            ).alignment = align_center

        # Lebar kolom
        ws_keperluan.column_dimensions["A"].width = 28
        ws_keperluan.column_dimensions["B"].width = 20

        data_jenis.to_excel(
            writer,
            index=False,
            sheet_name="Jenis Tamu"
        )

        ws_jenis = writer.sheets["Jenis Tamu"]

        # Sisipkan 2 baris untuk judul
        ws_jenis.insert_rows(1, amount=2)

        # Judul periode
        ws_jenis.merge_cells("A1:B1")
        ws_jenis["A1"] = judul_periode
        ws_jenis["A1"].font = Font(
            bold=True,
            size=14
        )
        ws_jenis["A1"].alignment = align_center

        # Header
        for cell in ws_jenis[3]:
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = align_center
            cell.border = border_tabel

        # Isi tabel
        for row in ws_jenis.iter_rows(
            min_row=4,
            max_row=ws_jenis.max_row,
            min_col=1,
            max_col=2
        ):
            for cell in row:
                cell.border = border_tabel
                cell.alignment = Alignment(
                    vertical="center"
                )

        # Kolom jumlah dibuat center
        for row in range(4, ws_jenis.max_row + 1):
            ws_jenis.cell(
                row=row,
                column=2
            ).alignment = align_center

        # Lebar kolom
        ws_jenis.column_dimensions["A"].width = 30
        ws_jenis.column_dimensions["B"].width = 14

        df_mentah.to_excel(
            writer,
            index=False,
            sheet_name="Data Mentah"
        )

        ws_mentah = writer.sheets["Data Mentah"]

        # Freeze header
        ws_mentah.freeze_panes = "A2"

        # Autofilter
        ws_mentah.auto_filter.ref = ws_mentah.dimensions

        # Style header
        for cell in ws_mentah[1]:
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = align_center
            cell.border = border_tabel

        # Style isi tabel
        for row in ws_mentah.iter_rows(
            min_row=2,
            max_row=ws_mentah.max_row,
            min_col=1,
            max_col=ws_mentah.max_column
        ):
            for cell in row:
                cell.border = border_tabel
                cell.alignment = Alignment(
                    vertical="center"
                )

        # Mapping nama header ke nomor kolom
        header_map = {
            cell.value: cell.column
            for cell in ws_mentah[1]
        }


        # Format tanggal
        if "Tanggal Kehadiran" in header_map:
            kolom_tanggal = header_map["Tanggal Kehadiran"]

            for row in range(2, ws_mentah.max_row + 1):
                ws_mentah.cell(
                    row=row,
                    column=kolom_tanggal
                ).number_format = "dd/mm/yyyy"


        # Format jam
        for nama_kolom in ["Jam Datang", "Jam Pulang"]:
            if nama_kolom in header_map:
                nomor_kolom = header_map[nama_kolom]

                for row in range(2, ws_mentah.max_row + 1):
                    ws_mentah.cell(
                        row=row,
                        column=nomor_kolom
                    ).number_format = "hh:mm"

        # Lebar kolom
        lebar_kolom = {
            "Timestamp": 22,
            "Nama": 24,
            "Jenis Tamu": 30,
            "Keterangan Tamu": 30,
            "Keperluan": 24,
            "Detail Keperluan": 36,
            "Tanggal Kehadiran": 18,
            "Jam Datang": 14,
            "Jam Pulang": 14,
        }


        for nama_kolom, lebar in lebar_kolom.items():
            if nama_kolom in header_map:
                huruf_kolom = ws_mentah.cell(
                    row=1,
                    column=header_map[nama_kolom]
                ).column_letter

                ws_mentah.column_dimensions[
                    huruf_kolom
                ].width = lebar

        if rekap_bulanan is not None:

            rekap_bulanan.to_excel(
                writer,
                index=False,
                header=False,
                sheet_name="Rekap Bulanan",
                startrow=4
            )

            ws_rekap = writer.sheets["Rekap Bulanan"]

            # =========================
            # JUDUL PERIODE
            # =========================
            ws_rekap.merge_cells("A1:G1")
            ws_rekap["A1"] = judul_periode

            ws_rekap["A1"].font = Font(
                bold=True,
                size=14
            )

            ws_rekap["A1"].alignment = align_center

            # =========================
            # HEADER BERTINGKAT
            # =========================
            ws_rekap.merge_cells("A3:A4")
            ws_rekap.merge_cells("B3:F3")
            ws_rekap.merge_cells("G3:G4")

            ws_rekap["A3"] = "Bulan"
            ws_rekap["B3"] = "Keperluan"
            ws_rekap["G3"] = "Jumlah"

            ws_rekap["B4"] = "Konsultasi Data"
            ws_rekap["C4"] = "Tugas Mitra"
            ws_rekap["D4"] = "Kepentingan Dinas"
            ws_rekap["E4"] = "Bertemu Orang"
            ws_rekap["F4"] = "Lainnya"

            # Style header
            for row in ws_rekap.iter_rows(
                min_row=3,
                max_row=4,
                min_col=1,
                max_col=7
            ):
                for cell in row:
                    cell.fill = fill_header
                    cell.font = font_header
                    cell.alignment = align_center
                    cell.border = border_tabel

            # =========================
            # ISI TABEL
            # =========================
            for row in ws_rekap.iter_rows(
                min_row=5,
                max_row=ws_rekap.max_row,
                min_col=1,
                max_col=7
            ):
                for cell in row:
                    cell.border = border_tabel
                    cell.alignment = align_center

            # Bulan rata kiri
            for row in range(5, ws_rekap.max_row + 1):
                ws_rekap.cell(
                    row=row,
                    column=1
                ).alignment = Alignment(
                    horizontal="left",
                    vertical="center"
                )

            # =========================
            # LEBAR KOLOM
            # =========================
            ws_rekap.column_dimensions["A"].width = 14
            ws_rekap.column_dimensions["B"].width = 20
            ws_rekap.column_dimensions["C"].width = 16
            ws_rekap.column_dimensions["D"].width = 20
            ws_rekap.column_dimensions["E"].width = 18
            ws_rekap.column_dimensions["F"].width = 14
            ws_rekap.column_dimensions["G"].width = 12

    return output.getvalue()

def siapkan_data_excel(df_periode):
    # Total tamu
    total_tamu = len(df_periode)

    # Keperluan
    data_kunjungan = (
        df_periode["Keperluan"]
        .value_counts()
        .reindex(
            kategori_keperluan,
            fill_value=0
        )
        .rename_axis("Keperluan Kunjungan")
        .reset_index(name="Jumlah Kunjungan")
    )

    # Jenis tamu
    data_jenis = (
        df_periode["Jenis Tamu"]
        .value_counts()
        .reindex(
            kategori_jenis_tamu,
            fill_value=0
        )
        .rename_axis("Jenis Tamu")
        .reset_index(name="Jumlah")
    )

    # Keperluan terbanyak
    if not data_kunjungan.empty:
        keperluan_terbanyak = (
            data_kunjungan
            .sort_values("Jumlah Kunjungan", ascending=False)
            .iloc[0]["Keperluan Kunjungan"]
        )
    else:
        keperluan_terbanyak = "-"

    # Durasi
    tanggal_string = df_periode[
        "Tanggal Kehadiran"
    ].dt.strftime("%Y-%m-%d")

    waktu_datang = pd.to_datetime(
        tanggal_string + " " + df_periode["Jam Datang"].astype(str),
        errors="coerce"
    )

    waktu_pulang = pd.to_datetime(
        tanggal_string + " " + df_periode["Jam Pulang"].astype(str),
        errors="coerce"
    )

    durasi = (
        waktu_pulang - waktu_datang
    ).dt.total_seconds() / 60

    durasi = durasi[durasi >= 0]

    if not durasi.empty:
        nilai_durasi = f"{durasi.mean():.0f} Menit"
    else:
        nilai_durasi = "-"

    # Ringkasan
    df_ringkasan = pd.DataFrame({
        "Indikator": [
            "Total Tamu",
            "Keperluan Terbanyak",
            "Rata-Rata Waktu Kunjungan"
        ],
        "Nilai": [
            f"{total_tamu} Orang",
            keperluan_terbanyak,
            nilai_durasi
        ]
    })

    return (
        df_ringkasan,
        data_kunjungan,
        data_jenis
    )

def get_triwulan(bulan):
    if bulan <= 3:
        return 1
    elif bulan <= 6:
        return 2
    elif bulan <= 9:
        return 3
    else:
        return 4

NAMA_BULAN = [
    "",
    "Januari", "Februari", "Maret", "April",
    "Mei", "Juni", "Juli", "Agustus",
    "September", "Oktober", "November", "Desember"
]

def siapkan_filter_periode(df, kolom_tanggal):
    df_temp = df.copy()

    df_temp["_tanggal_filter"] = pd.to_datetime(
        df_temp[kolom_tanggal],
        format="%d/%m/%Y",
        errors="coerce"
    )

    df_temp = df_temp.dropna(
        subset=["_tanggal_filter"]
    )

    periode_bulan = (
        df_temp["_tanggal_filter"]
        .dt.to_period("M")
        .drop_duplicates()
        .sort_values(ascending=False)
        .tolist()
    )

    opsi_periode = []

    # Bulanan
    for periode in periode_bulan:
        opsi_periode.append(
            f"{NAMA_BULAN[periode.month]} {periode.year}"
        )

    # Triwulanan
    triwulan_tersedia = set()

    for periode in periode_bulan:
        triwulan_tersedia.add(
            (
                periode.year,
                get_triwulan(periode.month)
            )
        )

    for tahun, triwulan in sorted(
        triwulan_tersedia,
        reverse=True
    ):
        opsi_periode.append(
            f"Triwulan {romawi[triwulan]} {tahun}"
        )

    return df_temp, opsi_periode

def filter_data_periode(df, periode_pilih):

    if periode_pilih.startswith("Triwulan"):

        _, nomor_romawi, tahun = periode_pilih.split()

        nomor_triwulan = romawi_ke_angka[nomor_romawi]
        tahun = int(tahun)

        bulan_awal = (nomor_triwulan - 1) * 3 + 1
        bulan_akhir = bulan_awal + 2

        hasil = df[
            (df["_tanggal_filter"].dt.year == tahun) &
            (df["_tanggal_filter"].dt.month >= bulan_awal) &
            (df["_tanggal_filter"].dt.month <= bulan_akhir)
        ].copy()

        nama_file = (
            f"Triwulan_{nomor_triwulan}_{tahun}"
        )

    else:

        nama_bulan, tahun = periode_pilih.split()
        tahun = int(tahun)

        nomor_bulan = NAMA_BULAN.index(nama_bulan)

        hasil = df[
            (df["_tanggal_filter"].dt.year == tahun) &
            (df["_tanggal_filter"].dt.month == nomor_bulan)
        ].copy()

        nama_file = (
            f"{nama_bulan}_{tahun}"
        )

    hasil = hasil.drop(
        columns=["_tanggal_filter"],
        errors="ignore"
    )

    return hasil, nama_file

def buat_excel_filter(df, nama_sheet="Data"):
    output = BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        df.to_excel(
            writer,
            index=False,
            sheet_name=nama_sheet
        )

        ws = writer.sheets[nama_sheet]

        # =========================
        # STYLE DASAR
        # =========================
        fill_header = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        font_header = Font(
            bold=True,
            color="FFFFFF"
        )

        garis_tipis = Side(
            style="thin",
            color="B7B7B7"
        )

        border_tabel = Border(
            left=garis_tipis,
            right=garis_tipis,
            top=garis_tipis,
            bottom=garis_tipis
        )

        # =========================
        # HEADER
        # =========================
        for cell in ws[1]:
            cell.fill = fill_header
            cell.font = font_header
            cell.border = border_tabel
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

        # =========================
        # ISI TABEL
        # =========================
        for row in ws.iter_rows(
            min_row=2,
            max_row=ws.max_row,
            min_col=1,
            max_col=ws.max_column
        ):
            for cell in row:
                cell.border = border_tabel
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

        # Freeze header
        ws.freeze_panes = "A2"

        # Autofilter
        ws.auto_filter.ref = ws.dimensions

        # =========================
        # LEBAR KOLOM OTOMATIS
        # =========================
        for column_cells in ws.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter

            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            ws.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                45
            )

    return output.getvalue()

romawi = {
    1: "I",
    2: "II",
    3: "III",
    4: "IV"
}

romawi_ke_angka = {
    "I": 1,
    "II": 2,
    "III": 3,
    "IV": 4
}

target_skd = {
    2026: {
        1: 7,
        2: 8,
        3: 7,
        4: 8,
    }
}

target_tahunan_skd = 30

laporan_skd = {
    2026: {
        1: "https://link-laporan-triwulan-1",
        2: "https://link-laporan-triwulan-2",
        3: None,
        4: None,
    }
}

def buat_rekap_bulanan(
    df_periode,
    nama_bulan,
    tambah_total=False,
    bulan_wajib=None
):

    df_temp = df_periode.copy()

    df_temp["Bulan"] = (
        df_temp["Tanggal Kehadiran"]
        .dt.month
        .map(lambda x: nama_bulan[x])
    )

    rekap = pd.crosstab(
        df_temp["Bulan"],
        df_temp["Keperluan"]
    )

    # Pastikan bulan yang tidak punya data tetap muncul
    if bulan_wajib is not None:
        rekap = rekap.reindex(
            bulan_wajib,
            fill_value=0
        )

    # Pastikan semua kategori selalu ada
    for kategori in kategori_keperluan:
        if kategori not in rekap.columns:
            rekap[kategori] = 0

    rekap = rekap[kategori_keperluan]

    rekap["Jumlah"] = rekap.sum(axis=1)

    rekap = rekap.reset_index()

    urutan_bulan = {
        nama_bulan[i]: i
        for i in range(1, 13)
    }

    rekap["_urutan"] = rekap["Bulan"].map(urutan_bulan)

    rekap = (
        rekap
        .sort_values("_urutan")
        .drop(columns="_urutan")
        .reset_index(drop=True)
    )

    # Tambahkan baris total
    if tambah_total:
        baris_total = {
            "Bulan": "Jumlah"
        }

        for kolom in kategori_keperluan:
            baris_total[kolom] = rekap[kolom].sum()

        baris_total["Jumlah"] = rekap["Jumlah"].sum()

        rekap = pd.concat(
            [
                rekap,
                pd.DataFrame([baris_total])
            ],
            ignore_index=True
        )

    return rekap
    
class HasilPengaduan(BaseModel):
    tanggal: str
    platform: str
    topik: str
    saran_pengaduan: str
    status_tindak_lanjut: str
    tindak_lanjut: str

class HasilPemutakhiran(BaseModel):
    tanggal_update: str
    kanal_digital: str
    jenis_konten: str
    topik_konten: str

class HasilFAQ(BaseModel):
    tanggal: str
    platform: str
    topik: str
    faq_key: str
    faq_judul: str
    pertanyaan: str
    status_tindak_lanjut: str
    tindak_lanjut: str

class RespondenSKD(BaseModel):
    tanggal_cacah: str
    nama: str
    status_kuesioner: str

class HasilAnalisisSKD(BaseModel):
    responden: list[RespondenSKD]

def analisis_screenshot_pengaduan(files):
    client = genai.Client(
        api_key=st.secrets["GEMINI_API_KEY"]
    )

    contents = []

    # Masukkan semua screenshot sesuai urutan drag
    for file in files:
        contents.append(
            types.Part.from_bytes(
                data=file.getvalue(),
                mime_type=file.type
            )
        )

    # Prompt ditaruh setelah seluruh gambar
    contents.append(
        """
        Seluruh gambar di atas adalah rangkaian screenshot
        dari SATU saran/pengaduan yang sama dan sudah diberikan
        dalam urutan percakapan yang benar.

        Baca seluruh screenshot sebagai satu kesatuan konteks.

        Ekstrak informasi berikut:

        1. tanggal
           Tanggal utama saran/pengaduan jika terlihat.
           Gunakan format DD/MM/YYYY.
           Jika tidak dapat dipastikan, isi "-".

        2. platform
           Contoh: Instagram, WhatsApp, Email, Facebook,
           LAPOR!, atau platform lain yang terlihat.

        3. topik
           Ringkas topik utama dalam frasa pendek.

        4. saran_pengaduan
           Ringkas inti saran, pertanyaan, apresiasi,
           atau pengaduan dari pengguna.

        5. status_tindak_lanjut
           Hanya boleh salah satu dari:
           "Selesai"
           "Belum Direspon"

           Gunakan "Selesai" jika terlihat sudah ada jawaban/
           tanggapan dari petugas.
           Gunakan "Belum Direspon" jika belum terlihat tanggapan.

        6. tindak_lanjut
           Ringkas jawaban/tanggapan petugas jika ada.
           Jika tidak ada, isi "(Belum Ada Balasan)".

        Jangan mengarang informasi yang tidak terlihat.
        """
    )

    response = client.models.generate_content(
        model="gemini-3.6-flash",
        contents=contents,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=HasilPengaduan
        )
    )

    return response.parsed

def analisis_screenshot_pemutakhiran(files):
    client = genai.Client(
        api_key=st.secrets["GEMINI_API_KEY"]
    )

    contents = []

    for file in files:
        contents.append(
            types.Part.from_bytes(
                data=file.getvalue(),
                mime_type=file.type
            )
        )

    tahun_sekarang = datetime.now().year
    contents.append(
        f"""
        Seluruh gambar di atas adalah bukti dari SATU pemutakhiran
        kanal digital yang sama.

        Baca seluruh screenshot sebagai satu kesatuan konteks.

        Ekstrak:

        1. tanggal_update
           Ambil tanggal publikasi/pemutakhiran yang terlihat pada screenshot.
           Gunakan format DD/MM/YYYY.
           Aturan tahun:
           - Jika tanggal, bulan, dan tahun terlihat jelas, gunakan tahun yang terlihat.
           - Jika hanya tanggal dan bulan yang terlihat tetapi tahun tidak ditampilkan,
             gunakan tahun saat ini: {tahun_sekarang}.
           - Jangan menebak atau mengarang tahun lain.
           - Jika tanggal dan bulan tidak dapat dipastikan, isi "-".

        2. kanal_digital
           Nama kanal atau platform yang diperbarui.
           Contoh: Website BPS, Instagram, Portal PPID,
           Facebook, YouTube, atau kanal lain yang terlihat.

        3. jenis_konten
           Identifikasi jenis/format konten yang terlihat.
           Contoh:
           - Website: Tabel, Publikasi, Infografis, Berita, Halaman Web
           - Instagram: Reels, Feed Post, Carousel, Story, Infografis
           - YouTube: Video, Shorts
           - Kanal lain: gunakan jenis konten yang paling sesuai dengan yang terlihat.
           Jika tidak dapat dipastikan, isi "-".
           - Untuk Instagram, gunakan aturan:
             Feed Post = unggahan feed satu gambar/foto biasa.
             Carousel = unggahan feed yang terdiri dari beberapa slide/gambar.
             Infografis = unggahan visual informatif yang harus menyajikan data, statistik,
             fakta, atau informasi dalam bentuk grafis.
             Reels = konten video Reels.
             Story = konten Instagram Story.

        4. topik_konten
            Identifikasi topik utama/substansi dari konten yang diperbarui.
            - Baca dan gabungkan informasi yang terlihat pada konten,
              termasuk teks pada gambar/video dan caption.
            - Utamakan informasi yang paling spesifik menjelaskan
              kegiatan, informasi, data, atau materi yang dipublikasikan.
            - Jangan hanya menyalin caption jika terdapat informasi
              yang lebih substantif pada gambar/video.
            - Jika terdapat nama kegiatan/program yang jelas,
              pertahankan nama tersebut dalam hasil.
            - Tulis sebagai judul/ringkasan singkat dan informatif.
            - Jangan mengarang informasi yang tidak terlihat.

        Jangan mengarang informasi yang tidak terlihat.
        """
    )

    response = client.models.generate_content(
        model="gemini-3.6-flash",
        contents=contents,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=HasilPemutakhiran
        )
    )

    return response.parsed

def analisis_screenshot_faq(files):
    client = genai.Client(
        api_key=st.secrets["GEMINI_API_KEY"]
    )

    contents = []

    for file in files:
        contents.append(
            types.Part.from_bytes(
                data=file.getvalue(),
                mime_type=file.type
            )
        )


    tahun_sekarang = datetime.now().year
    contents.append(
        f"""
        Seluruh gambar di atas adalah rangkaian screenshot
        dari SATU pertanyaan publik yang sama dan sudah diberikan
        dalam urutan percakapan yang benar.

        Baca seluruh screenshot sebagai satu kesatuan konteks.

        Ekstrak:

        1. tanggal
           Ambil tanggal awal pertanyaan yang terlihat pada screenshot.
           Gunakan format DD/MM/YYYY.

           Aturan tahun:
           - Jika tanggal, bulan, dan tahun terlihat jelas, gunakan tahun yang terlihat.
           - Jika hanya tanggal dan bulan yang terlihat tetapi tahun tidak ditampilkan,
             gunakan tahun saat ini: {tahun_sekarang}.
           - Jangan menebak atau mengarang tahun lain yang tidak terlihat.
           - Jika tanggal dan bulan tidak dapat dipastikan, isi "-".

        2. platform
           Platform tempat pertanyaan diterima.
           Contoh: Instagram, WhatsApp, Email, Facebook,
           LAPOR!, atau platform lain yang terlihat.

        3. topik
           Ringkas kategori/topik pertanyaan dalam frasa pendek.

        4. pertanyaan
           Ringkas pertanyaan utama dari pengguna.
           Pertahankan inti informasi yang ditanyakan.

        5. status_tindak_lanjut
           Hanya boleh salah satu:
           "Selesai"
           "Belum Direspon"

           Gunakan "Selesai" jika terlihat sudah ada jawaban
           atau tanggapan dari petugas.

           Gunakan "Belum Direspon" jika belum ada jawaban.

        6. tindak_lanjut
           Ringkas jawaban/tanggapan petugas jika ada.
           Jika belum ada tanggapan, isi "(Belum Ada Balasan)".

        7. faq_key
           Buat label kanonik yang sangat singkat untuk mengelompokkan
           pertanyaan dengan maksud yang sama.

           Gunakan inti kebutuhan pengguna, bukan cara kalimat ditulis.

           Contoh:
           "Apakah tersedia data umur tunggal?"
           "Saya membutuhkan data penduduk menurut umur tunggal"
           "Dimana mencari data umur tunggal Kepulauan Seribu?"
           semuanya harus menghasilkan faq_key yang sama:
           "data umur tunggal"

           Gunakan huruf kecil dan jangan sertakan kata umum seperti
           "permintaan", "menanyakan", "cara bertanya", atau nama platform.
        
        8. faq_judul
           Buat judul FAQ yang singkat, natural, dan berbentuk pertanyaan.

           Judul harus mewakili inti pertanyaan pengguna,
           bukan menyalin kalimat panjang percakapan.

           Contoh:
           "Pengguna menanyakan ketersediaan data umur tunggal
           Kepulauan Seribu tahun 2010-2020"
           menjadi:
           "Apakah tersedia data umur tunggal Kepulauan Seribu?"

           Gunakan bahasa Indonesia yang ringkas dan jelas.

        Jangan mengarang informasi yang tidak terlihat.
        """
    )

    response = client.models.generate_content(
        model="gemini-3.6-flash",
        contents=contents,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=HasilFAQ
        )
    )

    return response.parsed

def analisis_screenshot_skd(files):
    client = genai.Client(
        api_key=st.secrets["GEMINI_API_KEY"]
    )

    contents = []

    for file in files:
        contents.append(
            types.Part.from_bytes(
                data=file.getvalue(),
                mime_type=file.type
            )
        )

    contents.append(
        """
        Seluruh gambar di atas adalah screenshot progres Survei
        Kebutuhan Data (SKD).

        Screenshot dapat berasal dari:
        1. Rekap Responden Terverifikasi
        2. Rekap Calon Responden

        Ekstrak SEMUA responden yang terlihat pada seluruh screenshot.

        Untuk setiap responden, hasilkan:

        tanggal_cacah
        - Ambil tanggal cacah yang terlihat.
        - Gunakan format DD/MM/YYYY.
        - Jika tidak dapat dipastikan, isi "-".

        nama
        - Ambil nama responden persis seperti terlihat.
        - Jangan membuat nama baru.

        status_kuesioner
        - Jika responden berada pada daftar Responden Terverifikasi,
          isi persis:
          "Sudah terverifikasi"

        - Jika berada pada daftar Calon Responden, baca kondisi
          Blok 1 sampai Blok 4.

          Interpretasi simbol:
          centang biru = blok sudah terisi dan sudah diverifikasi
          silang merah = blok sudah terisi tetapi belum diverifikasi
          tanda "-" = blok belum diisi
          nilai seperti "0/1" = isian pada blok belum lengkap

        Buat kesimpulan status yang singkat dan jelas.

        Contoh:
        "Belum terisi lengkap pada Blok 3"
        "Blok 4 belum diisi"
        "Sudah terisi lengkap tapi belum diverifikasi"
        "Belum seluruh blok diverifikasi"

        Jangan mengarang responden atau status yang tidak terlihat.

        Jika nama yang sama terlihat lebih dari sekali pada screenshot,
        keluarkan satu record dengan kondisi yang paling mutakhir /
        paling lengkap.
        """
    )

    response = client.models.generate_content(
        model="gemini-3.6-flash",
        contents=contents,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=HasilAnalisisSKD
        )
    )

    return response.parsed

# --- TAB UTAMA ---
tab1, tab2, tab3, tab4, tab5 = st.tabs([
    "Overview",
    "Saran & Pengaduan",
    "Pertanyaan & FAQ",
    "Progres SKD",
    "Kanal Digital"
])

# ==========================================
# TAB 1: OVERVIEW KUNJUNGAN PST
# ==========================================
with tab1:
    from services.google_sheets import get_google_sheet_data

    @st.fragment(run_every="1h") # Auto-update setiap 1 jam
    def overview_kunjungan_pst():

        # Ambil Google Sheets
        df = get_google_sheet_data()

        if df.empty:
            st.warning("Belum ada data kunjungan.")
            return

        # Ambil waktu update
        waktu_update = datetime.now(
            ZoneInfo("Asia/Jakarta")
        ).strftime("%d/%m/%Y %H:%M WIB")

        # --- PREPROCESSING DATA ---
        df["Tanggal Kehadiran"] = pd.to_datetime(
            df["Tanggal Kehadiran"],
            errors="coerce"
        )

        # Buang baris yang tanggalnya tidak valid / NA
        df = df.dropna(subset=["Tanggal Kehadiran"])

        if df.empty:
            st.warning("Data tanggal kunjungan belum valid.")
            return

        nama_bulan = [
            "",
            "Januari", "Februari", "Maret", "April",
            "Mei", "Juni", "Juli", "Agustus",
            "September", "Oktober", "November", "Desember"
        ]

        periode_bulan_tersedia = (
            df["Tanggal Kehadiran"]
            .dt.to_period("M")
            .drop_duplicates()
            .sort_values(ascending=False)
            .tolist()
        )

        opsi_periode = []

        for periode in periode_bulan_tersedia:
            opsi_periode.append(
                f"{nama_bulan[periode.month]} {periode.year}"
            )

        # Opsi Triwulan
        triwulan_tersedia = set()

        for periode in periode_bulan_tersedia:
            nomor_triwulan = get_triwulan(periode.month)

            triwulan_tersedia.add(
                (periode.year, nomor_triwulan)
            )

        for tahun, triwulan in sorted(
            triwulan_tersedia,
            reverse=True
        ):
            opsi_periode.append(
                f"Triwulan {romawi[triwulan]} {tahun}"
            )

        # =========================
        # PILIH PERIODE DASHBOARD
        # =========================

        # Placeholder judul supaya tetap tampil di atas baris kontrol
        judul_dashboard = st.empty()

        col_update, col_periode, col_download = st.columns(
            [7, 1.8, 0.65],
            gap="small",
            vertical_alignment="center"
        )

        with col_update:
            st.caption(
                f"🗘 *Data diperbarui otomatis setiap 1 jam* | "
                f"*Terakhir diperbarui: {waktu_update}*"
            )

        with col_periode:
            periode_pilih = st.selectbox(
                "Periode",
                opsi_periode,
                index=0,
                label_visibility="collapsed"
            )


        # =========================
        # FILTER PERIODE TERPILIH
        # =========================

        is_triwulan = periode_pilih.startswith("Triwulan")

        if is_triwulan:

            _, nomor_triwulan, tahun_periode = periode_pilih.split()

            nomor_triwulan = romawi_ke_angka[nomor_triwulan]
            tahun_periode = int(tahun_periode)

            bulan_awal = (nomor_triwulan - 1) * 3 + 1
            bulan_akhir = bulan_awal + 2

            bulan_triwulan = [
                nama_bulan[i]
                for i in range(bulan_awal, bulan_akhir + 1)
            ]

            # Data periode yang sedang ditampilkan
            df_periode = df[
                (df["Tanggal Kehadiran"].dt.year == tahun_periode) &
                (df["Tanggal Kehadiran"].dt.month >= bulan_awal) &
                (df["Tanggal Kehadiran"].dt.month <= bulan_akhir)
            ].copy()

            # Triwulan sebelumnya untuk delta
            if nomor_triwulan > 1:
                triwulan_lalu = nomor_triwulan - 1
                tahun_lalu = tahun_periode
            else:
                triwulan_lalu = 4
                tahun_lalu = tahun_periode - 1

            bulan_awal_lalu = (triwulan_lalu - 1) * 3 + 1
            bulan_akhir_lalu = bulan_awal_lalu + 2

            df_periode_lalu = df[
                (df["Tanggal Kehadiran"].dt.year == tahun_lalu) &
                (df["Tanggal Kehadiran"].dt.month >= bulan_awal_lalu) &
                (df["Tanggal Kehadiran"].dt.month <= bulan_akhir_lalu)
            ].copy()

            label_periode = (
                f"Triwulan {romawi[nomor_triwulan]} "
                f"{tahun_periode}"
            )

            nama_file_periode = (
                f"Triwulan_{nomor_triwulan}_{tahun_periode}"
            )

            label_perbandingan = "dari triwulan lalu"

        else:

            bulan_triwulan = None
            periode_terpilih = None

            for periode in periode_bulan_tersedia:
                label = (
                    f"{nama_bulan[periode.month]} {periode.year}"
                )

                if label == periode_pilih:
                    periode_terpilih = periode
                    break

            # Data bulan yang sedang ditampilkan
            df_periode = df[
                (df["Tanggal Kehadiran"].dt.year == periode_terpilih.year) &
                (df["Tanggal Kehadiran"].dt.month == periode_terpilih.month)
            ].copy()

            # Bulan sebelumnya untuk delta
            periode_lalu = periode_terpilih - 1

            df_periode_lalu = df[
                (df["Tanggal Kehadiran"].dt.year == periode_lalu.year) &
                (df["Tanggal Kehadiran"].dt.month == periode_lalu.month)
            ].copy()

            label_periode = (
                f"{nama_bulan[periode_terpilih.month]} "
                f"{periode_terpilih.year}"
            )

            nama_file_periode = (
                f"{nama_bulan[periode_terpilih.month]}_"
                f"{periode_terpilih.year}"
            )

            label_perbandingan = "dari bulan lalu"

        # Judul dashboard mengikuti periode
        judul_dashboard.markdown(
            f'<div class="jala-page-title">Ringkasan Kunjungan PST {label_periode}</div>',
            unsafe_allow_html=True
        )

        # =========================
        # 1. TOTAL KUNJUNGAN
        # =========================

        total_tamu = len(df_periode)
        total_periode_lalu = len(df_periode_lalu)

        if total_periode_lalu > 0:
            perubahan_tamu = (
                (total_tamu - total_periode_lalu)
                / total_periode_lalu
            ) * 100

            delta_total = (
                f"{perubahan_tamu:+.0f}% "
                f"{label_perbandingan}"
            )
        else:
            delta_total = None


        # =========================
        # 2. KEPERLUAN TERBANYAK
        # =========================

        distribusi_keperluan = (
            df_periode["Keperluan"]
            .dropna()
            .value_counts()
        )

        if not distribusi_keperluan.empty:
            keperluan_terbanyak = distribusi_keperluan.idxmax()
            jumlah_terbanyak = distribusi_keperluan.max()

            persen_terbanyak = (
                jumlah_terbanyak / total_tamu * 100
                if total_tamu > 0
                else 0
            )

            delta_keperluan = (
                f"{persen_terbanyak:.0f}% dari total"
            )

        else:
            keperluan_terbanyak = "-"
            delta_keperluan = None


        # =========================
        # 3. DURASI KUNJUNGAN
        # =========================

        tanggal_string = (
            df_periode["Tanggal Kehadiran"]
            .dt.strftime("%Y-%m-%d")
        )

        waktu_datang = pd.to_datetime(
            tanggal_string + " " +
            df_periode["Jam Datang"].astype(str),
            errors="coerce"
        )

        waktu_pulang = pd.to_datetime(
            tanggal_string + " " +
            df_periode["Jam Pulang"].astype(str),
            errors="coerce"
        )

        durasi = (
            waktu_pulang - waktu_datang
        ).dt.total_seconds() / 60

        durasi = durasi[durasi >= 0]

        if not durasi.empty:
            rata_rata_durasi = durasi.mean()
            nilai_durasi = f"{rata_rata_durasi:.0f} Menit"
        else:
            rata_rata_durasi = None
            nilai_durasi = "-"


        # Durasi periode sebelumnya
        delta_durasi = None

        if not df_periode_lalu.empty:

            tanggal_lalu = (
                df_periode_lalu["Tanggal Kehadiran"]
                .dt.strftime("%Y-%m-%d")
            )

            datang_lalu = pd.to_datetime(
                tanggal_lalu + " " +
                df_periode_lalu["Jam Datang"].astype(str),
                errors="coerce"
            )

            pulang_lalu = pd.to_datetime(
                tanggal_lalu + " " +
                df_periode_lalu["Jam Pulang"].astype(str),
                errors="coerce"
            )

            durasi_lalu = (
                pulang_lalu - datang_lalu
            ).dt.total_seconds() / 60

            durasi_lalu = durasi_lalu[
                durasi_lalu >= 0
            ]

            if (
                not durasi_lalu.empty
                and rata_rata_durasi is not None
            ):
                selisih = (
                    rata_rata_durasi -
                    durasi_lalu.mean()
                )

                delta_durasi = f"{selisih:+.0f} Menit"


        # =========================
        # DATA BAR CHART
        # =========================

        data_kunjungan = (
            df_periode["Keperluan"]
            .dropna()
            .value_counts()
            .reset_index()
        )

        data_kunjungan.columns = [
            "Keperluan Kunjungan",
            "Jumlah Kunjungan"
        ]

        data_kunjungan = (
            data_kunjungan
            .sort_values("Jumlah Kunjungan")
        )


        # =========================
        # DATA DONUT CHART
        # =========================

        data_jenis = (
            df_periode["Jenis Tamu"]
            .dropna()
            .value_counts()
            .reset_index()
        )

        data_jenis.columns = [
            "Jenis Tamu",
            "Jumlah"
        ]


        # =========================
        # DATA EXCEL
        # =========================

        (
            df_ringkasan_download,
            data_kunjungan_download,
            data_jenis_download
        ) = siapkan_data_excel(df_periode)

        rekap_bulanan = buat_rekap_bulanan(
            df_periode,
            nama_bulan,
            tambah_total=is_triwulan,
            bulan_wajib=bulan_triwulan
        )

        excel_overview = buat_excel_overview(
            df_ringkasan_download,
            data_kunjungan_download,
            data_jenis_download,
            df_periode,
            rekap_bulanan,
            nama_file_periode
        )

        # Tombol selalu aktif
        with col_download:
            st.download_button(
                label="⎙ Unduh Data",
                data=excel_overview,
                file_name=(
                    f"Overview_Kunjungan_Tamu_"
                    f"{nama_file_periode}.xlsx"
                ),
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                use_container_width=True
            )

        # =========================
        # METRIC CARDS
        # =========================
        col1, col2, col3 = st.columns(3)

        col1.metric(
            "Total Kunjungan",
            f"{total_tamu} Kunjungan",
            delta=delta_total
        )

        col2.metric(
            "Keperluan Terbanyak",
            keperluan_terbanyak,
            delta=delta_keperluan
        )

        col3.metric(
            "Rata-Rata Waktu Kunjungan",
            nilai_durasi,
            delta=delta_durasi,
            delta_color="inverse"
        )
            
        # =========================
        # CHART
        # =========================
        col_grafik, col_info = st.columns([65, 35])

        PALETTE_BIRU = [
            "#C6E7FF",
            "#64B5F6",
            "#1E88E5",
            "#0D47A1"
        ]

        with col_grafik:

            with st.container(border=True):

                fig = px.bar(
                    data_kunjungan,
                    x="Jumlah Kunjungan",
                    y="Keperluan Kunjungan",
                    orientation="h",
                    text="Jumlah Kunjungan",
                    title=(
                        "Distribusi Keperluan "
                        "Kunjungan Tamu PST"
                    ),
                    color="Jumlah Kunjungan",
                    color_continuous_scale=PALETTE_BIRU
                )

                fig.update_traces(
                    textposition="outside",
                    textfont=dict(size=13),
                    cliponaxis=False
                )

                fig.update_layout(
                    height=250,
                    margin=dict(
                        t=30,
                        b=10,
                        l=10,
                        r=30
                    ),
                    coloraxis_showscale=False,
                    xaxis_title="Jumlah Kunjungan",
                    yaxis_title="",
                    plot_bgcolor="rgba(0,0,0,0)"
                )

                st.plotly_chart(
                    fig,
                    use_container_width=True
                )

        with col_info:

            with st.container(border=True):

                fig_donut = px.pie(
                    data_jenis,
                    names="Jenis Tamu",
                    values="Jumlah",
                    hole=0.5,
                    title="Proporsi Jenis Tamu",
                    color_discrete_sequence=PALETTE_BIRU
                )

                fig_donut.update_traces(
                    textinfo="none",
                    selector=dict(type="pie"),
                    domain=dict(
                        x=[0.08, 0.92],
                        y=[0.20, 1.00]
                    ),
                    hovertemplate=(
                        "<b>%{label}</b><br>"
                        "%{value} kunjungan<br>"
                        "%{percent}"
                        "<extra></extra>"
                    )
                )

                fig_donut.update_layout(
                    showlegend=True,
                    legend=dict(
                        orientation="h",
                        yanchor="top",
                        y=0.10,  
                        xanchor="center",
                        x=0.5,
                        font=dict(size=10),
                        entrywidth=160,
                        entrywidthmode="pixels"
                    ),
                    height=250,
                    margin=dict(
                        t=30,
                        b=5,
                        l=75,
                        r=70
                    ),
                    uniformtext_minsize=9,
                    uniformtext_mode="show"
                )

                st.plotly_chart(
                    fig_donut,
                    use_container_width=True
                )

    overview_kunjungan_pst()

# ==========================================
# TAB 2: SARAN & PENGADUAN 
# ==========================================
with tab2:
    st.markdown(
        '<div class="jala-page-title">Repositori Saran & Pengaduan</div>',
        unsafe_allow_html=True
    )

    if st.session_state.pop("sp_simpan_sukses", False):
        st.success("✅ Data berhasil disimpan!")

    # Upload Box Screenshot
    with st.expander("⛶ Upload Screenshot", expanded=True):

        # Untuk reset file uploader setelah pengaduan berhasil disimpan
        if "sp_uploader_version" not in st.session_state:
            st.session_state["sp_uploader_version"] = 0

        uploaded_files = st.file_uploader(
            "Unggah satu atau beberapa screenshot untuk 1 saran/pengaduan. Jika percakapan terdiri dari beberapa screenshot, unggah seluruh screenshot sekaligus sesuai urutan percakapan.",
            type=["jpg", "jpeg", "png"],
            accept_multiple_files=True,
            key=f"sp_uploader_{st.session_state['sp_uploader_version']}"
        )
        
        # Cek apakah list uploaded_files ada isinya (bukan pengecekan None)
        if uploaded_files:

            # Kalau lebih dari 1 screenshot, tampilkan fitur pengurutan
            if len(uploaded_files) > 1:

                st.markdown("###### ⇄ Urutan Screenshot")
                st.caption(
                    "Drag nama file untuk menyesuaikan urutan percakapan "
                    "sebelum dianalisis."
                )

                label_to_file = {
                    f"{idx + 1}. {file.name}": file
                    for idx, file in enumerate(uploaded_files)
                }

                # Key berubah kalau daftar file berubah
                file_signature = "_".join(
                    sorted(file.name for file in uploaded_files)
                )

                urutan_file = sort_items(
                    list(label_to_file.keys()),
                    direction="vertical",
                    key=f"sp_sorter_{file_signature}",
                    custom_style="""
                    .sortable-component,
                    .sortable-container,
                    ul,
                    ol {
                        background: transparent !important;
                        border: none !important;
                        box-shadow: none !important;
                        padding: 0 !important;
                        margin: 0 !important;
                    }

                    .sortable-item {
                        background: #F3F8FE !important;
                        color: #16477C !important;

                        border: 1px solid #D6E4F2 !important;
                        border-radius: 8px !important;

                        min-height: 30px !important;
                        padding: 2px 7px !important;
                        margin-bottom: 3px !important;

                        font-size: 13px !important;
                        font-weight: 500 !important;

                        box-shadow: none !important;
                    }
                    """
                )

                ordered_files = [
                    label_to_file[label]
                    for label in urutan_file
                    if label in label_to_file
                ]

            # Kalau cuma 1 screenshot, tidak perlu sorter
            else:
                ordered_files = uploaded_files


            # PREVIEW
            st.markdown("###### 👁 Preview Screenshot")

            jumlah_kolom = min(len(ordered_files), 4)

            kolom_preview = st.columns(
                [1, 1, 1, 1],
                gap="small"
            )

            for idx, file in enumerate(ordered_files):
                kolom = kolom_preview[idx % 4]

                with kolom:
                    st.image(
                        file,
                        caption=f"Screenshot {idx + 1}",
                        use_container_width=True
                    )

            # TOMBOL ANALISIS
            if st.button(
                "Analisis Screenshot",
                key="btn_analisis_sp"
            ):

                try:
                    with st.spinner("Menganalisis data"):
                        hasil = analisis_screenshot_pengaduan(
                            ordered_files
                        )

                    st.session_state["hasil_analisis_sp"] = (
                        hasil.model_dump()
                    )

                except Exception as e:
                    st.error(
                        f"Gagal menganalisis screenshot: {e}"
                    )

            # Form hasil analisis
            if "hasil_analisis_sp" in st.session_state:

                hasil = st.session_state["hasil_analisis_sp"]

                st.markdown("### Hasil Analisis")

                tanggal_review = st.text_input(
                    "Tanggal",
                    value=hasil["tanggal"],
                    key="review_tanggal"
                )

                platform_review = st.text_input(
                    "Platform",
                    value=hasil["platform"],
                    key="review_platform"
                )

                topik_review = st.text_input(
                    "Topik",
                    value=hasil["topik"],
                    key="review_topik"
                )

                saran_review = st.text_area(
                    "Saran/Pengaduan",
                    value=hasil["saran_pengaduan"],
                    key="review_saran"
                )

                pilihan_status = [
                    "⌛︎ Belum Direspon",
                    "✔ Selesai"
                ]

                status_ai = hasil["status_tindak_lanjut"]

                index_status = (
                    1 if status_ai == "Selesai"
                    else 0
                )

                status_review = st.selectbox(
                    "Status Tindak Lanjut",
                    pilihan_status,
                    index=index_status,
                    key="review_status"
                )

                tindak_lanjut_review = st.text_area(
                    "Tindak Lanjut",
                    value=hasil["tindak_lanjut"],
                    key="review_tindak_lanjut"
                )

                if st.button(
                    "Simpan Pengaduan",
                    key="btn_simpan_pengaduan"
                ):
                    try:
                        with st.spinner("Menyimpan data..."):

                            # 1. Buat nama folder bukti
                            nama_folder = (
                                f"SP_{tanggal_review.replace('/', '-')}_"
                                f"{platform_review.replace(' ', '-')}"
                            )

                            # 2. Upload screenshot ke Google Drive
                            hasil_drive = upload_screenshot(
                                ordered_files,
                                nama_folder,
                                st.secrets["DRIVE_SP_FOLDER_ID"]
                            )

                            folder_link = hasil_drive["folder_link"]

                            # 3. Rapikan status sebelum masuk Google Sheets
                            status_sheet = (
                                "✔ Selesai"
                                if "Selesai" in status_review
                                else "⌛︎ Belum Direspon"
                            )

                            # 4. Simpan hasil review ke Database JALA-SERIBU
                            append_pengaduan_row([
                                tanggal_review,
                                platform_review,
                                topik_review,
                                saran_review,
                                status_sheet,
                                tindak_lanjut_review,
                                folder_link
                            ])

                            # Tandai bahwa penyimpanan berhasil
                            st.session_state["sp_simpan_sukses"] = True
                            # Hapus hasil analisis
                            st.session_state.pop("hasil_analisis_sp", None)
                            # Reset file uploader
                            st.session_state["sp_uploader_version"] += 1
                            # Jalankan ulang halaman dalam keadaan bersih
                            st.rerun()

                    except Exception as e:
                        st.error(f"Gagal menyimpan pengaduan: {e}")

    st.write("##### 🗐 Tabel Rekapitulasi Pengaduan")
    
    # Tabel rekapitulasi Saran dan Pengaduan
    df_pengaduan = get_pengaduan_data()

    if df_pengaduan.empty:
        st.info("Belum ada data saran/pengaduan yang tersimpan")

    else:
        df_pengaduan_filter, opsi_pengaduan = (
            siapkan_filter_periode(
                df_pengaduan,
                "Tanggal"
            )
        )

        col_info_sp, col_periode_sp, col_download_sp = st.columns(
            [7, 1.8, 0.65],
            gap="small",
            vertical_alignment="center"
        )

        with col_info_sp:
            st.caption("Filter data berdasarkan periode")

        with col_periode_sp:
            periode_sp = st.selectbox(
                "Periode Pengaduan",
                opsi_pengaduan,
                index=0,
                label_visibility="collapsed",
                key="periode_pengaduan"
            )

        df_pengaduan_tampil, nama_file_sp = filter_data_periode(
            df_pengaduan_filter,
            periode_sp
        )

        excel_pengaduan = buat_excel_filter(
            df_pengaduan_tampil,
            "Saran dan Pengaduan"
        )

        with col_download_sp:
            st.download_button(
                "⎙ Unduh Data",
                data=excel_pengaduan,
                file_name=f"Pengaduan_{nama_file_sp}.xlsx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                key="download_pengaduan"
            )

        st.dataframe(
            df_pengaduan_tampil,
            use_container_width=True,
            hide_index=True,
            row_height=100,
            column_config={
                "Tanggal": st.column_config.TextColumn(
                    "Tanggal",
                    width="small",
                ),
                "Platform": st.column_config.TextColumn(
                    "Platform",
                    width="small",
                ),
                "Topik": st.column_config.TextColumn(
                    "Topik",
                    width="medium",
                ),
                "Saran/Pengaduan": st.column_config.TextColumn(
                    "Saran/Pengaduan",
                    width="medium",
                ),
                "Status Tindak Lanjut": st.column_config.TextColumn(
                    "Status Tindak Lanjut",
                    width="small",
                ),
                "Tindak Lanjut": st.column_config.TextColumn(
                    "Tindak Lanjut",
                    width="large",
                ),
                "Bukti": st.column_config.LinkColumn(
                    "Bukti",
                    display_text="🖿 Lihat Bukti",
                    width="small",
                ),
            },
        )

# ==========================================
# TAB 3: PERTANYAAN DAN FAQ
# ==========================================
with tab3:
    st.markdown(
        '<div class="jala-page-title">Repositori Pertanyaan Publik dan FAQ</div>',
        unsafe_allow_html=True
    )

    if st.session_state.pop("faq_simpan_sukses", False):
        st.success("✅ Pertanyaan berhasil disimpan!")

    if "faq_uploader_version" not in st.session_state:
        st.session_state["faq_uploader_version"] = 0

    # 1. FORM INPUT SCREENSHOT PERTANYAAN BARU
    with st.expander("⛶ Upload Screenshot", expanded=False):
            uploaded_faqs = st.file_uploader(
                "Unggah satu atau beberapa screenshot untuk 1 pertanyaan. Jika percakapan terdiri dari beberapa screenshot, unggah seluruh screenshot sekaligus sesuai urutan percakapan.", 
                type=["jpg", "jpeg", "png"], 
                accept_multiple_files=True, 
                key=f"faq_uploader_{st.session_state['faq_uploader_version']}"
            )
            
            # Pengecekan apakah list uploaded_faqs ada isinya
            if uploaded_faqs:

                # Kalau screenshot lebih dari 1, tampilkan sorter
                if len(uploaded_faqs) > 1:

                    st.markdown("###### ⇄ Urutan Screenshot")
                    st.caption(
                        "Drag nama file untuk menyesuaikan urutan "
                        "percakapan sebelum dianalisis."
                    )

                    label_to_file_faq = {
                        f"{idx + 1}. {file.name}": file
                        for idx, file in enumerate(uploaded_faqs)
                    }

                    file_signature_faq = "_".join(
                        sorted(file.name for file in uploaded_faqs)
                    )

                    urutan_faq = sort_items(
                        list(label_to_file_faq.keys()),
                        direction="vertical",
                        key=f"faq_sorter_{file_signature_faq}",
                        custom_style="""
                        .sortable-component,
                        .sortable-container,
                        ul,
                        ol {
                            background: transparent !important;
                            border: none !important;
                            box-shadow: none !important;
                            padding: 0 !important;
                            margin: 0 !important;
                        }

                        .sortable-item {
                            background: #F3F8FE !important;
                            color: #16477C !important;
                            border: 1px solid #D6E4F2 !important;
                            border-radius: 8px !important;
                            min-height: 30px !important;
                            padding: 2px 7px !important;
                            margin-bottom: 3px !important;
                            font-size: 13px !important;
                            font-weight: 500 !important;
                            box-shadow: none !important;
                        }
                        """
                    )

                    ordered_faqs = [
                        label_to_file_faq[label]
                        for label in urutan_faq
                        if label in label_to_file_faq
                    ]

                else:
                    ordered_faqs = uploaded_faqs


                # Preview
                st.markdown("###### 👁 Preview Screenshot")

                kolom_preview_faq = st.columns(
                    [1, 1, 1, 1],
                    gap="small"
                )

                for idx, file in enumerate(ordered_faqs):
                    kolom = kolom_preview_faq[idx % 4]

                    with kolom:
                        st.image(
                            file,
                            caption=f"Screenshot {idx + 1}",
                            use_container_width=True
                        )

                # Analisis
                if st.button(
                    "Analisis Screenshot",
                    key="btn_analisis_faq"
                ):
                    try:
                        with st.spinner("Menganalisis pertanyaan..."):

                            hasil_faq = analisis_screenshot_faq(
                                ordered_faqs
                            )

                        st.session_state["hasil_analisis_faq"] = (
                            hasil_faq.model_dump()
                        )

                    except Exception as e:
                        st.error(
                            f"Gagal menganalisis screenshot: {e}"
                        )

                # Review
                if "hasil_analisis_faq" in st.session_state:
                    hasil_faq = st.session_state["hasil_analisis_faq"]

                    # FAQ Key untuk pengelompokan pertanyaan serupa
                    faq_key_faq_review = hasil_faq["faq_key"]

                    st.markdown("### Hasil Analisis")

                    tanggal_faq_review = st.text_input(
                        "Tanggal",
                        value=hasil_faq["tanggal"],
                        key="review_tanggal_faq"
                    )

                    platform_faq_review = st.text_input(
                        "Platform",
                        value=hasil_faq["platform"],
                        key="review_platform_faq"
                    )

                    topik_faq_review = st.text_input(
                        "Topik",
                        value=hasil_faq["topik"],
                        key="review_topik_faq"
                    )

                    pertanyaan_faq_review = st.text_area(
                        "Pertanyaan",
                        value=hasil_faq["pertanyaan"],
                        key="review_pertanyaan_faq"
                    )

                    pilihan_status_faq = [
                        "⌛︎ Belum Direspon",
                        "✔ Selesai"
                    ]

                    status_ai_faq = hasil_faq["status_tindak_lanjut"]

                    index_status_faq = (
                        1 if status_ai_faq == "Selesai"
                        else 0
                    )

                    status_faq_review = st.selectbox(
                        "Status Tindak Lanjut",
                        pilihan_status_faq,
                        index=index_status_faq,
                        key="review_status_faq"
                    )

                    tindak_lanjut_faq_review = st.text_area(
                        "Tindak Lanjut",
                        value=hasil_faq["tindak_lanjut"],
                        key="review_tindak_lanjut_faq"
                    )

                    faq_key_faq_review = hasil_faq["faq_key"]

                    if st.button(
                        "Simpan Pertanyaan",
                        key="btn_simpan_faq"
                    ):
                        try:
                            with st.spinner("Menyimpan pertanyaan..."):

                                # 1. Nama folder bukti
                                nama_folder_faq = (
                                    f"FAQ_{tanggal_faq_review.replace('/', '-')}_"
                                    f"{platform_faq_review.replace(' ', '-')}"
                                )

                                # 2. Upload screenshot ke folder Pertanyaan & FAQ
                                hasil_drive_faq = upload_screenshot(
                                    ordered_faqs,
                                    nama_folder_faq,
                                    st.secrets["DRIVE_FAQ_FOLDER_ID"]
                                )

                                folder_link_faq = hasil_drive_faq["folder_link"]

                                # 3. Rapikan status untuk Google Sheets
                                status_sheet_faq = (
                                    "✔ Selesai"
                                    if "Selesai" in status_faq_review
                                    else "⌛︎ Belum Direspon"
                                )

                                # 4. Simpan ke Database JALA-SERIBU
                                append_faq_row([
                                    tanggal_faq_review,
                                    platform_faq_review,
                                    topik_faq_review,
                                    faq_key_faq_review,
                                    pertanyaan_faq_review,
                                    status_sheet_faq,
                                    tindak_lanjut_faq_review,
                                    folder_link_faq
                                ])

                            st.session_state["faq_simpan_sukses"] = True
                            st.session_state.pop("hasil_analisis_faq", None)
                            st.session_state["faq_uploader_version"] += 1

                            st.rerun()

                        except Exception as e:
                            st.error(
                                f"Gagal menyimpan pertanyaan: {e}"
                            )

    # 2. TABEL REKAPITULASI SEMUA PERTANYAAN
    st.write("##### 🗐 Tabel Rekapitulasi Pertanyaan Masuk")

    df_faq = get_faq_data()

    if df_faq.empty:
        st.info("Belum ada data pertanyaan yang tersimpan.")

    else:
        # =========================
        # FILTER PERIODE TABEL
        # =========================
        df_faq_filter, opsi_faq = siapkan_filter_periode(
            df_faq,
            "Tanggal"
        )

        col_info_faq, col_periode_faq, col_download_faq = st.columns(
            [7, 1.8, 0.65],
            gap="small",
            vertical_alignment="center"
        )

        with col_info_faq:
            st.caption("Filter data berdasarkan periode")

        with col_periode_faq:
            periode_faq = st.selectbox(
                "Periode Pertanyaan",
                opsi_faq,
                index=0,
                label_visibility="collapsed",
                key="periode_pertanyaan"
            )

        df_faq_tampil, nama_file_faq = filter_data_periode(
            df_faq_filter,
            periode_faq
        )

        # FAQ Key tidak perlu ditampilkan / diekspor
        df_faq_display = df_faq_tampil.drop(
            columns=["FAQ Key"],
            errors="ignore"
        )

        # =========================
        # EXCEL SESUAI FILTER
        # =========================
        excel_faq = buat_excel_filter(
            df_faq_display,
            "Pertanyaan Masuk"
        )

        with col_download_faq:
            st.download_button(
                "⎙ Unduh Data",
                data=excel_faq,
                file_name=f"Pertanyaan_{nama_file_faq}.xlsx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                key="download_pertanyaan"
            )

        # =========================
        # TABEL SESUAI FILTER
        # =========================
        st.dataframe(
            df_faq_display,
            use_container_width=True,
            hide_index=True,
            row_height=100,
            column_config={
                "Tanggal": st.column_config.TextColumn(
                    "Tanggal",
                    width="small",
                ),
                "Platform": st.column_config.TextColumn(
                    "Platform",
                    width="small",
                ),
                "Topik": st.column_config.TextColumn(
                    "Topik",
                    width="medium",
                ),
                "Pertanyaan": st.column_config.TextColumn(
                    "Pertanyaan",
                    width="medium",
                ),
                "Status Tindak Lanjut": st.column_config.TextColumn(
                    "Status Tindak Lanjut",
                    width="small",
                ),
                "Tindak Lanjut": st.column_config.TextColumn(
                    "Tindak Lanjut",
                    width="large",
                ),
                "Bukti": st.column_config.LinkColumn(
                    "Bukti",
                    display_text="🖿 Buka Bukti",
                    width="small",
                ),
            },
        )

    # Frekuensi FAQ memakai seluruh data, bukan hasil filter
    frekuensi_topik = (
        df_faq["FAQ Key"]
        .fillna("-")
        .value_counts()
        if not df_faq.empty
        else pd.Series(dtype=int)
    )

    # 3. PERTANYAAN POPULER / TOP FAQ (KARTU RINGKASAN JAWABAN)
    st.write("##### ✦ Pertanyaan Paling Populer (FAQ)")

    if df_faq.empty:
        st.info("Belum ada data untuk membentuk FAQ.")

    else:
        # Ambil maksimal 3 topik yang paling sering ditanyakan
        top_faq = frekuensi_topik.head(3)

        for posisi, (faq_key, jumlah) in enumerate(top_faq.items()):

            kelompok_faq = df_faq[
                df_faq["FAQ Key"].fillna("-") == faq_key
            ].copy()

            if kelompok_faq.empty:
                continue

            # Ubah tanggal supaya bisa pilih record terbaru
            kelompok_faq["_tanggal_dt"] = pd.to_datetime(
                kelompok_faq["Tanggal"],
                format="%d/%m/%Y",
                errors="coerce"
            )

            kelompok_faq = kelompok_faq.sort_values(
                "_tanggal_dt",
                ascending=False
            )

            # Ambil pertanyaan terbaru sebagai judul FAQ
            pertanyaan_populer = kelompok_faq.iloc[0]["Pertanyaan"]

            # Cari jawaban/tindak lanjut yang statusnya sudah selesai
            jawaban_selesai = kelompok_faq[
                kelompok_faq["Status Tindak Lanjut"]
                .astype(str)
                .str.contains("Selesai", case=False, na=False)
            ]

            if not jawaban_selesai.empty:
                jawaban_resmi = jawaban_selesai.iloc[0]["Tindak Lanjut"]
            else:
                jawaban_resmi = "Belum ada jawaban resmi untuk pertanyaan ini."

            with st.expander(
                f"❓ {pertanyaan_populer} (Ditanyakan {jumlah}x)",
                expanded=(posisi == 0)
            ):
                st.write("**Jawaban Resmi:**")
                st.info(jawaban_resmi)
    
# ==========================================
# TAB 4: PROGRES SKD
# ==========================================
with tab4:
    sekarang = datetime.now()

    tahun_skd = sekarang.year
    triwulan_skd = get_triwulan(sekarang.month)
    nama_triwulan_skd = romawi[triwulan_skd]

    target_triwulan_skd = target_skd.get(
        tahun_skd, {}
    ).get(
        triwulan_skd, 0
    )

    # Ambil data progres SKD dari Google Sheets
    df_skd = get_skd_data()

    jumlah_triwulan_skd = 0
    jumlah_tahunan_skd = 0

    if not df_skd.empty:

        df_skd_hitung = df_skd.copy()

        # Ubah Tanggal Cacah menjadi datetime
        df_skd_hitung["_tanggal"] = pd.to_datetime(
            df_skd_hitung["Tanggal Cacah"],
            format="%d/%m/%Y",
            errors="coerce"
        )

        # Ambil hanya responden yang sudah terverifikasi
        df_terverifikasi = df_skd_hitung[
            df_skd_hitung["Status Kuesioner"]
            .astype(str)
            .str.strip()
            .str.lower()
            .eq("sudah terverifikasi")
        ].copy()

        # Jumlah terverifikasi selama tahun aktif
        jumlah_tahunan_skd = len(
            df_terverifikasi[
                df_terverifikasi["_tanggal"].dt.year == tahun_skd
            ]
        )

        # Jumlah terverifikasi pada triwulan aktif
        df_triwulan_aktif = df_terverifikasi[
            (df_terverifikasi["_tanggal"].dt.year == tahun_skd)
            &
            (df_terverifikasi["_tanggal"].dt.quarter == triwulan_skd)
        ]

        jumlah_triwulan_skd = len(df_triwulan_aktif)

    persen_triwulan_skd = (
        jumlah_triwulan_skd / target_triwulan_skd
        if target_triwulan_skd > 0
        else 0
    )

    persen_tahunan_skd = (
        jumlah_tahunan_skd / target_tahunan_skd
        if target_tahunan_skd > 0
        else 0
    )

    progress_triwulan_skd = min(
        persen_triwulan_skd,
        1.0
    )

    progress_tahunan_skd = min(
        persen_tahunan_skd,
        1.0
    )

    st.markdown(
        f"""
        <div class="jala-page-title">
            Progres Survei Kebutuhan Data (SKD) Triwulan {nama_triwulan_skd} {tahun_skd}
        </div>
        """,
        unsafe_allow_html=True
    )

    if "skd_simpan_sukses" in st.session_state:
        hasil_simpan = st.session_state.pop(
            "skd_simpan_sukses"
        )

        st.success(
            f"✅ Progres berhasil disimpan! "
            f"{hasil_simpan['inserted']} data baru, "
            f"{hasil_simpan['updated']} data diperbarui."
        )

    col_skd1, col_skd2 = st.columns([3, 1])
    
    with col_skd1:
        sub_col_skd1, sub_col_skd2 = st.columns([1, 1])
        
        with sub_col_skd1:
            st.write(
                f"##### Jumlah Responden Triwulan "
                f"{nama_triwulan_skd} {tahun_skd}"
            )
            st.progress(
                progress_triwulan_skd,
                text=(
                    f"{persen_triwulan_skd * 100:.2f}% dari Target "
                    f"({jumlah_triwulan_skd}/{target_triwulan_skd} Responden)"
                ).replace(".", ",")
            )

        with sub_col_skd2:
            st.write("##### Jumlah Responden Dalam Setahun")
            st.progress(
                progress_tahunan_skd,
                text=(
                    f"{persen_tahunan_skd * 100:.2f}% dari Target "
                    f"({jumlah_tahunan_skd}/{target_tahunan_skd} Responden)"
                ).replace(".", ",")
            )
        
        # FORM INPUT SCREENSHOT
        if "skd_uploader_version" not in st.session_state:
            st.session_state["skd_uploader_version"] = 0

        with st.expander("⛶ Upload Screenshot", expanded=True):
                uploaded_skds = st.file_uploader(
                    "Upload screenshot progres pengisian survei. Dapat upload multiple files.", 
                    type=["jpg", "jpeg", "png"], 
                    accept_multiple_files=True, 
                    key=f"skd_uploader_{st.session_state['skd_uploader_version']}"
                )
                
                # Pengecekan apakah ada file dalam list uploaded_skds
                if uploaded_skds:
                    st.markdown("###### 👁 Preview Screenshot")

                    kolom_preview_skd = st.columns(
                        [1, 1, 1, 1],
                        gap="small"
                    )

                    for idx, file in enumerate(uploaded_skds):
                        kolom = kolom_preview_skd[idx % 4]

                        with kolom:
                            st.image(
                                file,
                                caption=f"Screenshot {idx + 1}",
                                use_container_width=True
                            )

                    if st.button(
                        "Analisis Screenshot",
                        key="btn_analisis_skd"
                    ):
                        try:
                            with st.spinner("Menganalisis progres SKD..."):
                                hasil_skd = analisis_screenshot_skd(
                                    uploaded_skds
                                )

                            st.session_state["hasil_analisis_skd"] = [
                                item.model_dump()
                                for item in hasil_skd.responden
                            ]

                        except Exception as e:
                            st.error(
                                f"Gagal menganalisis screenshot SKD: {e}"
                            )

                    if "hasil_analisis_skd" in st.session_state:

                        st.markdown("### Hasil Analisis")

                        df_review_skd = pd.DataFrame(
                            st.session_state["hasil_analisis_skd"]
                        )

                        st.dataframe(
                            df_review_skd,
                            use_container_width=True,
                            hide_index=True
                        )

                        # TOMBOL MASIH DI DALAM IF
                        if st.button(
                            "Simpan Progres",
                            key="btn_simpan_skd"
                        ):
                            try:
                                with st.spinner("Menyimpan progres SKD..."):

                                    inserted = 0
                                    updated = 0

                                    for item in st.session_state["hasil_analisis_skd"]:
                                        hasil = upsert_skd_row(
                                            item["tanggal_cacah"],
                                            item["nama"],
                                            item["status_kuesioner"]
                                        )

                                        if hasil == "inserted":
                                            inserted += 1
                                        elif hasil == "updated":
                                            updated += 1

                                st.session_state["skd_simpan_sukses"] = {
                                    "inserted": inserted,
                                    "updated": updated
                                }

                                st.session_state.pop(
                                    "hasil_analisis_skd",
                                    None
                                )

                                st.session_state["skd_uploader_version"] += 1

                                st.rerun()

                            except Exception as e:
                                st.error(f"Gagal menyimpan progres SKD: {e}")

                        

        # TABEL PROGRES PENGISIAN SURVEI
        df_skd = get_skd_data()

        if df_skd.empty:
            st.info("Belum ada data progres SKD")
        else:
            st.dataframe(
                df_skd,
                use_container_width=True,
                hide_index=True,
            )
                    
    with col_skd2:
        st.write("##### 🗁 Akses Laporan Resmi SKD")

        laporan_tahun = laporan_skd.get(tahun_skd, {})

        for tw in range(1, 5):
            nama_tw = romawi[tw]
            link_laporan = laporan_tahun.get(tw)

            if link_laporan:
                st.link_button(
                    f"🗎 Laporan SKD Triwulan {nama_tw} {tahun_skd}",
                    link_laporan,
                    use_container_width=True
                )
            else:
                st.button(
                    f"🗎 Laporan SKD Triwulan {nama_tw} {tahun_skd} (Drafting)",
                    disabled=True,
                    use_container_width=True,
                    key=f"laporan_skd_{tahun_skd}_{tw}"
                )

# ==========================================
# TAB 5: PEMUTAKHIRAN KANAL DIGITAL
# ==========================================
with tab5:
    st.markdown(
        '<div class="jala-page-title">Pemutakhiran Kanal Digital</div>',
        unsafe_allow_html=True
    )

    if st.session_state.pop("web_simpan_sukses", False):
        st.success("✅ Data berhasil disimpan!")

    # Ambil data asli
    df_website = get_pemutakhiran_data()

    if df_website.empty:
        st.info("Belum ada data pemutakhiran kanal digital yang tersimpan.")

        df_website_tampil = df_website

        total_update = 0
        kanal_aktif = "-"
        jumlah_kanal_aktif = 0
        update_terakhir = "-"
        tanggal_update_terakhir = None

    else:
        # =========================
        # FILTER PERIODE
        # =========================
        df_website_filter, opsi_web = siapkan_filter_periode(
            df_website,
            "Tanggal Update"
        )

        col_info_web, col_periode_web, col_download_web = st.columns(
            [7, 1.8, 0.65],
            gap="small",
            vertical_alignment="center"
        )

        with col_info_web:
            st.caption("Filter data berdasarkan periode")

        with col_periode_web:
            periode_web = st.selectbox(
                "Periode Kanal Digital",
                opsi_web,
                index=0,
                label_visibility="collapsed",
                key="periode_web"
            )

        df_website_tampil, nama_file_web = filter_data_periode(
            df_website_filter,
            periode_web
        )

        # =========================
        # DOWNLOAD SESUAI FILTER
        # =========================
        excel_web = buat_excel_filter(
            df_website_tampil,
            "Kanal Digital"
        )

        with col_download_web:
            st.download_button(
                "⎙ Unduh Data",
                data=excel_web,
                file_name=f"Kanal_Digital_{nama_file_web}.xlsx",
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                key="download_web"
            )

        # =========================
        # KPI SESUAI FILTER
        # =========================
        total_update = len(df_website_tampil)

        if df_website_tampil.empty:
            kanal_aktif = "-"
            jumlah_kanal_aktif = 0
            update_terakhir = "-"
            tanggal_update_terakhir = None

        else:
            kanal_counts = (
                df_website_tampil["Kanal Digital"]
                .value_counts()
            )

            kanal_aktif = kanal_counts.index[0]
            jumlah_kanal_aktif = int(kanal_counts.iloc[0])

            tanggal_series = pd.to_datetime(
                df_website_tampil["Tanggal Update"],
                format="%d/%m/%Y",
                errors="coerce"
            )

            tanggal_max = tanggal_series.max()

            if pd.notna(tanggal_max):
                tanggal_sekarang = pd.Timestamp.today().normalize()

                selisih_hari = (
                    tanggal_sekarang -
                    tanggal_max.normalize()
                ).days

                if selisih_hari == 0:
                    update_terakhir = "Hari ini"
                elif selisih_hari == 1:
                    update_terakhir = "Kemarin"
                else:
                    update_terakhir = f"{selisih_hari} hari lalu"

                tanggal_update_terakhir = (
                    tanggal_max.strftime("%d/%m/%Y")
                )

            else:
                update_terakhir = "-"
                tanggal_update_terakhir = None


    # =========================
    # METRIC CARDS
    # =========================
    col_web1, col_web2, col_web3 = st.columns(3)

    col_web1.metric(
        "Total Update",
        f"{total_update} Konten"
    )

    col_web2.metric(
        "Kanal Paling Aktif",
        kanal_aktif,
        f"{jumlah_kanal_aktif} Update"
        if jumlah_kanal_aktif > 0
        else None
    )

    col_web3.metric(
        "Update Terakhir",
        update_terakhir,
        tanggal_update_terakhir,
        delta_arrow="off"
    )

    # 2. FORM INPUT SCREENSHOT UPDATE
    with st.expander("⛶ Upload Screenshot", expanded=True):
            if "web_uploader_version" not in st.session_state:
                st.session_state["web_uploader_version"] = 0

            uploaded_webs = st.file_uploader(
                "Unggah satu atau beberapa screenshot untuk 1 pemutakhiran kanal digital. Jika percakapan terdiri dari beberapa screenshot, unggah seluruh screenshot sekaligus sesuai urutan percakapan.", 
                type=["jpg", "jpeg", "png"], 
                accept_multiple_files=True, 
                key=f"web_uploader_{st.session_state['web_uploader_version']}"
            )
            
            # Pengecekan apakah ada file yang diunggah dalam list uploaded_webs
            if uploaded_webs:

                # Kalau screenshot lebih dari 1, tampilkan sorter
                if len(uploaded_webs) > 1:

                    st.markdown("###### ⇄ Urutan Screenshot")
                    st.caption(
                        "Drag nama file untuk menyesuaikan urutan bukti "
                        "sebelum dianalisis."
                    )

                    label_to_file_web = {
                        f"{idx + 1}. {file.name}": file
                        for idx, file in enumerate(uploaded_webs)
                    }

                    file_signature_web = "_".join(
                        sorted(file.name for file in uploaded_webs)
                    )

                    urutan_web = sort_items(
                        list(label_to_file_web.keys()),
                        direction="vertical",
                        key=f"web_sorter_{file_signature_web}",
                        custom_style="""
                        .sortable-component,
                        .sortable-container,
                        ul,
                        ol {
                            background: transparent !important;
                            border: none !important;
                            box-shadow: none !important;
                            padding: 0 !important;
                            margin: 0 !important;
                        }

                        .sortable-item {
                            background: #F3F8FE !important;
                            color: #16477C !important;
                            border: 1px solid #D6E4F2 !important;
                            border-radius: 8px !important;
                            min-height: 30px !important;
                            padding: 2px 7px !important;
                            margin-bottom: 3px !important;
                            font-size: 13px !important;
                            font-weight: 500 !important;
                            box-shadow: none !important;
                        }
                        """
                    )

                    ordered_webs = [
                        label_to_file_web[label]
                        for label in urutan_web
                        if label in label_to_file_web
                    ]

                else:
                    ordered_webs = uploaded_webs

                # Preview
                st.markdown("###### 👁 Preview Screenshot")

                kolom_preview_web = st.columns(
                    [1, 1, 1, 1],
                    gap="small"
                )

                for idx, file in enumerate(ordered_webs):
                    kolom = kolom_preview_web[idx % 4]

                    with kolom:
                        st.image(
                            file,
                            caption=f"Screenshot {idx + 1}",
                            use_container_width=True
                        )

                # Analisis
                if st.button(
                    "Analisis Screenshot",
                    key="btn_analisis_web"
                ):
                    try:
                        with st.spinner(
                            "Menganalisis data..."
                        ):
                            hasil_web = analisis_screenshot_pemutakhiran(
                                ordered_webs
                            )

                        st.session_state["hasil_analisis_web"] = (
                            hasil_web.model_dump()
                        )

                    except Exception as e:
                        st.error(
                            f"Gagal menganalisis screenshot: {e}"
                        )

                if "hasil_analisis_web" in st.session_state:

                    hasil_web = st.session_state["hasil_analisis_web"]

                    st.markdown("### Hasil Analisis")

                    tanggal_web_review = st.text_input(
                        "Tanggal Update",
                        value=hasil_web["tanggal_update"],
                        key="review_tanggal_web"
                    )

                    kanal_web_review = st.text_input(
                        "Kanal Digital",
                        value=hasil_web["kanal_digital"],
                        key="review_kanal_web"
                    )

                    jenis_web_review = st.text_input(
                        "Jenis Konten",
                        value=hasil_web["jenis_konten"],
                        key="review_jenis_web"
                    )

                    topik_web_review = st.text_area(
                        "Topik Konten",
                        value=hasil_web["topik_konten"],
                        key="review_topik_web"
                    )

                    if st.button(
                        "Simpan Pemutakhiran",
                        key="btn_simpan_web"
                    ):
                        try:
                            with st.spinner("Menyimpan pemutakhiran..."):

                                nama_folder = (
                                    f"UPDATE_{tanggal_web_review.replace('/', '-')}_"
                                    f"{kanal_web_review.replace(' ', '-')}"
                                )

                                hasil_drive_web = upload_screenshot(
                                    ordered_webs,
                                    nama_folder,
                                    st.secrets["DRIVE_WEB_FOLDER_ID"]
                                )

                                folder_link_web = hasil_drive_web["folder_link"]

                                append_pemutakhiran_row([
                                    tanggal_web_review,
                                    kanal_web_review,
                                    jenis_web_review,
                                    topik_web_review,
                                    folder_link_web
                                ])

                            st.session_state["web_simpan_sukses"] = True
                            st.session_state.pop("hasil_analisis_web", None)
                            st.session_state["web_uploader_version"] += 1

                            st.rerun()

                        except Exception as e:
                            st.error(
                                f"Gagal menyimpan pemutakhiran: {e}"
                            )

    # 3. TABEL REKAPITULASI LOG UPDATE WEBSITE
    st.write("##### 🗐 Tabel Rekapitulasi Log Update Kanal Digital")

    if df_website.empty:
        st.info("Belum ada data pemutakhiran kanal digital yang tersimpan.")

    elif df_website_tampil.empty:
        st.info("Tidak ada data pada periode yang dipilih.")

    else:
        st.dataframe(
            df_website_tampil,
            use_container_width=True,
            hide_index=True,
            column_config={
                "Bukti": st.column_config.LinkColumn(
                    "Bukti",
                    display_text="🖿 Buka Bukti"
                )
            }
        )

st.markdown(
    """
    <div class="jala-footer">
        © 2026 JALA-SERIBU · 
        <span class="footer-bps">BPS Kabupaten Kepulauan Seribu</span>
    </div>
    """,
    unsafe_allow_html=True
)